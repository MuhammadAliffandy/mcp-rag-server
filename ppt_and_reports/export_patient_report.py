import json
import sys
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_docx(json_path, out_path):
    if not os.path.exists(json_path):
        print(f"Error: {json_path} not found.")
        sys.exit(1)
        
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = Document()
    
    # Title
    title = doc.add_heading('ColonoSense Automated RAG Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Patient Info
    doc.add_heading('Patient Information', level=1)
    gt = data.get("ground_truth", {})
    
    pinfo = doc.add_paragraph()
    pinfo.add_run(f"Patient ID: ").bold = True
    pinfo.add_run(f"{data.get('patient_id')}\n")
    
    pinfo.add_run(f"Age / Sex: ").bold = True
    pinfo.add_run(f"{gt.get('Age', 'N/A')} / {gt.get('Sex', 'N/A')}\n")
    
    pinfo.add_run(f"Total Mayo Score: ").bold = True
    pinfo.add_run(f"{gt.get('total_mayo_score', 'N/A')}\n")
    
    pinfo.add_run(f"Endoscopic Score (MES): ").bold = True
    pinfo.add_run(f"{gt.get('mes_score', 'N/A')}\n")
    
    pinfo.add_run(f"Fecal Calprotectin: ").bold = True
    pinfo.add_run(f"{gt.get('fecal_calprotectin', 'N/A')}\n")
    
    pinfo.add_run(f"CRP: ").bold = True
    pinfo.add_run(f"{gt.get('crp', 'N/A')}\n")
    
    # RAG Answers
    doc.add_heading('Clinical Recommendations (Q1 - Q6)', level=1)
    
    responses = data.get("agent_responses", {})
    results = data.get("results", [])
    
    # Check if there are missing categories
    expected_categories = [
        "Q1.1", "Q1.2", "Q1.3",
        "Q2.1", "Q2.2", "Q2.3",
        "Q3.1", "Q3.2",
        "Q4.1", "Q4.2", "Q4.3",
        "Q5.1", "Q5.2", "Q5.3",
        "Q6.1", "Q6.2", "Q6.3"
    ]
    
    for category in expected_categories:
        cat_heading = doc.add_heading(f"{category}", level=2)
        
        ans = responses.get(category, "")
        if not ans or ans.strip() == "":
            ans = "No response generated. (This may be due to missing input data or system timeout)."
            p = doc.add_paragraph()
            run = p.add_run(ans)
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00) # Red
        else:
            doc.add_paragraph(ans.strip())
        
        # Add metrics
        res = next((r for r in results if r.get("category") == category), None)
        if res:
            p = doc.add_paragraph()
            r = p.add_run(f"Metrics: Data Acc {res.get('data_retrieval_acc', 0):.1f}% | Correctness {res.get('output_correctness', 0):.1f}%")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(out_path)
    print(f"Exported successfully to {out_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python export_patient_report.py <input.json> <output.docx>")
        sys.exit(1)
    export_docx(sys.argv[1], sys.argv[2])
